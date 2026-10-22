"""Package the LaTeX slides as a PowerPoint slideshow without changing layout."""
from pathlib import Path
import argparse

from pptx import Presentation
from pptx.util import Inches

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("rendered_pages", type=Path, help="PNG pages rendered from main.pdf")
args = parser.parse_args()
pages = sorted(args.rendered_pages.glob("page-*.png"))
assert len(pages) == 9, "Expected the nine reviewed slides."

deck = Presentation()
deck.slide_width, deck.slide_height = Inches(16), Inches(9)
deck.core_properties.title = "模型融合的后处理防御"
deck.core_properties.subject = "两种方案与攻击评估"
notes = [
    "汇报三件事：两套后处理防御及攻击评估。明确已完成、失败和仍未完成的研究，不宣称全部解决。",
    "原专家冻结。仅发布变换可优化，构造不读未来伙伴或测试反馈；目标任务和伙伴任务必须分开报告。",
    "偏置核空间保持是理想算术恒等式；原AGA没有处理关键最终偏置。高独立准确率和低融合准确率都不构成普遍安全证明。",
    "距离最大化目标确实存在。旧O/B上游仅开放任务SVD坐标，B输入范围较大而输出受限。有限探针不等于功能零空间。原优化为有限任务约束下的可行更新与非线性校正，不是未知伙伴安全博弈。",
    "上下两表是不同协议：两专家TA0.5与八专家TA0.2。后者只评三任务。O/B和A/B都有预算反例。新增功能保持对照通过新确认和完整MNIST任务门控，但TA99.34、AGA-target99.26、both99.31，保护无效，已关闭，未据结果调参。",
    "任务误差是教师对各竞争类别的相对间隔变化。O/B、A/B是512张训练审计，新对照是10000张历史测试，样本范围不同。新对照参数距离0.4999995、完整任务漂移0.0001205，保持通过；真实目标保护失败。有限SVD补空间幅度大不代表分类损害。",
    "Unmask整份参数减法0.3已完整评价，无保护控制也严重退化，不能据低分声称防御成功。区分完整官方代码、官方内核的native适配、论文公式实现与未取得代码。ProDistill需要校准图和任务归属，逐元素系数可有很高维度；它属于数据辅助融合。",
    "NUFILT跳过编辑位置，产生与无保护相同的合并模型，不等于恢复了原偏置。偏置清理的不同精度协议分列，旧数值失败保留。",
    "Defense2新的合规有效方法仍是最大缺口。下一步先给完整原生结构的可实现条件和可证伪假设；不要再机械放大噪声或扫描秩、种子和模块。",
]
for page, note in zip(pages, notes):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(page), 0, 0, width=deck.slide_width, height=deck.slide_height)
    slide.notes_slide.notes_text_frame.text = note + "\nLaTeX源稿：presentation/main.tex；数值由update_tables.py生成。"
output = Path(__file__).with_name("research_report.pptx")
deck.save(output)
print(output)

